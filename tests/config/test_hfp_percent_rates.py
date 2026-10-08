"""
Tests for reading the percent columns of the Debts and Fixed Assets sheets (#176).

Rates and commissions are percent numbers (4.5 for 4.5%) and are read as typed, however
small: a 0.5% real growth rate must not become 50%. Only a cell the workbook formats as a
percentage, which displays 4.50% and stores 0.045, is multiplied by 100.

Copyright (C) 2024-2026 Martin-D. Lacasse and The Owl Authors

This program is free software: you can redistribute it and/or modify
it under the terms of the GNU General Public License as published by
the Free Software Foundation, either version 3 of the License, or
(at your option) any later version.

This program is distributed in the hope that it will be useful,
but WITHOUT ANY WARRANTY; without even the implied warranty of
MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
GNU General Public License for more details.

You should have received a copy of the GNU General Public License
along with this program.  If not, see <https://www.gnu.org/licenses/>.
"""

import io
from datetime import date

import openpyxl
import pandas as pd
import pytest

import owlplanner as owl
from owlplanner.hfp_io import conditionDebtsAndFixedAssetsDF

thisyear = date.today().year


def _plan():
    return owl.Plan(["Joe", "Jane"], ["1964-03-15", "1964-09-15"], [89, 92], "rate", verbose=False)


def _tables(home_rate=0.5, loan_rate=0.9, commission=6.0):
    fixed = pd.DataFrame(
        [dict(active=True, name="home", type="residence", year=thisyear, basis=500000.0, value=800000.0,
              rate=home_rate, yod=0, commission=commission)]
    )
    debts = pd.DataFrame(
        [dict(active=True, name="mtg", type="mortgage", year=thisyear, term=30, amount=400000.0, rate=loan_rate)]
    )
    return conditionDebtsAndFixedAssetsDF(fixed, "Fixed Assets"), conditionDebtsAndFixedAssetsDF(debts, "Debts")


def _saved_workbook(tmp_path, **kwargs):
    p = _plan()
    p.houseLists["Fixed Assets"], p.houseLists["Debts"] = _tables(**kwargs)
    path = tmp_path / "HFP_rate.xlsx"
    p.saveHFP(str(path), overwrite=True)
    return path


def _read(source):
    q = _plan()
    q.readHFP(source)
    return q.houseLists["Fixed Assets"], q.houseLists["Debts"]


def _format_as_percent(path, cells):
    """Store each (sheet, column, value) as a fraction in a percent-formatted cell."""
    wb = openpyxl.load_workbook(path)
    for sheet, col, value in cells:
        ws = wb[sheet]
        k = [c.value for c in ws[1]].index(col) + 1
        cell = ws.cell(row=2, column=k)
        cell.value = value
        cell.number_format = "0.00%"
    wb.save(path)


def test_rates_below_one_survive_a_round_trip(tmp_path):
    fixed, debts = _read(str(_saved_workbook(tmp_path)))
    assert fixed["rate"].iloc[0] == pytest.approx(0.5)
    assert fixed["commission"].iloc[0] == pytest.approx(6.0)
    assert debts["rate"].iloc[0] == pytest.approx(0.9)


def test_percent_formatted_cells_are_read_as_percent(tmp_path):
    path = _saved_workbook(tmp_path, home_rate=0.0)
    _format_as_percent(path, [("Debts", "rate", 0.045), ("Fixed Assets", "commission", 0.06)])
    fixed, debts = _read(str(path))
    assert debts["rate"].iloc[0] == pytest.approx(4.5)
    assert fixed["commission"].iloc[0] == pytest.approx(6.0)
    assert fixed["rate"].iloc[0] == 0.0


def test_an_in_memory_upload_is_read(tmp_path):
    path = _saved_workbook(tmp_path)
    _format_as_percent(path, [("Debts", "rate", 0.045)])
    buf = io.BytesIO(path.read_bytes())
    fixed, debts = _read(buf)
    assert debts["rate"].iloc[0] == pytest.approx(4.5)
    assert fixed["rate"].iloc[0] == pytest.approx(0.5)


def test_a_blank_row_inside_the_table_keeps_cells_on_their_rows(tmp_path):
    """Percent cells below a blank row are converted on their own rows, and no other."""
    path = _saved_workbook(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Debts"]
    header = [c.value for c in ws[1]]
    first = [c.value for c in ws[2]]
    ws.insert_rows(2)  # blank row 2; the first loan moves to row 3
    second = dict(zip(header, first))
    second.update(name="car", type="loan", amount=20000.0, rate=0.07)
    for k, col in enumerate(header, start=1):
        ws.cell(row=4, column=k, value=second[col])
    ws.cell(row=4, column=header.index("rate") + 1).number_format = "0.00%"
    wb.save(path)

    _, debts = _read(str(path))
    rates = dict(zip(debts["name"], debts["rate"]))
    assert rates["mtg"] == pytest.approx(0.9)
    assert rates["car"] == pytest.approx(7.0)


def test_tables_given_as_dataframes_are_read_as_typed():
    q = _plan()
    fixed, debts = _tables()
    q.readHFP({"Fixed Assets": fixed, "Debts": debts, "Joe": pd.DataFrame({"year": [thisyear]}),
               "Jane": pd.DataFrame({"year": [thisyear]})})
    assert q.houseLists["Fixed Assets"]["rate"].iloc[0] == pytest.approx(0.5)
    assert q.houseLists["Debts"]["rate"].iloc[0] == pytest.approx(0.9)


def test_the_plan_uses_the_rate_as_entered(tmp_path):
    """The issue's repro: the home's final value and the loan payment survive a save and reload."""
    p = _plan()
    p.houseLists["Fixed Assets"], p.houseLists["Debts"] = _tables()
    path = tmp_path / "HFP_rate.xlsx"
    p.saveHFP(str(path), overwrite=True)
    q = _plan()
    q.readHFP(str(path))
    for plan in (p, q):
        plan.setRates("conservative")
        plan.processDebtsAndFixedAssets()
    assert q.fixed_assets_bequest_value == pytest.approx(p.fixed_assets_bequest_value)
    assert q.debt_payments_n[0] == pytest.approx(p.debt_payments_n[0])
