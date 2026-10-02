"""
The HFP download must carry the edits made in the Financial Profile tables.

The editors update the session's tables, while the workbook is written from the plan's, which
were refreshed only when the case ran. A ticked *Roth conv fixed* (or any other edit) made
since the last upload or run was therefore saved with its old value.

Copyright (C) 2024-2026 Martin-D. Lacasse and The Owl Authors

This program is free software: you can redistribute it and/or modify
it under the terms of the GNU General Public License as published by
the Free Software Foundation, either version 3 of the License, or
(at your option) any later version.

This program is distributed in the hope that it will be useful,
but WITHOUT ANY WARRANTY; without even the implied warranty of
MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
GNU General Public License for more details.

You should have received a copy of the GNU General Public License
along with this program.  If not, see <https://www.gnu.org/licenses/>.
"""

from datetime import date

import openpyxl
import pytest

import owlplanner as owl
import ui.owlbridge as owb


@pytest.mark.toml
def test_hfp_download_carries_table_edits(monkeypatch):
    plan = owl.readConfig("examples/Case_joe.toml", verbose=False)
    plan.mylog.setVerbose(False)
    name = plan.inames[0]

    # Edit the table the way the page does: a new DataFrame stored in the session.
    edited = plan.timeLists[name].copy()
    row = edited.index[edited["year"] == date.today().year + 2][0]
    assert not edited.at[row, "Roth conv fixed"]
    edited.at[row, "Roth conv fixed"] = True
    edited.at[row, "Roth conv"] = 12_345.0

    session = {
        "plan": plan,
        "iname0": name,
        "status": "single",
        "timeList0": edited,
        "hfpFileName": "HFP_joe.xlsx",
        "houseListDebts": plan.houseLists.get("Debts"),
        "houseListFixedAssets": plan.houseLists.get("Fixed Assets"),
    }
    monkeypatch.setattr(owb.kz, "getCaseKey", session.get)
    monkeypatch.setattr(owb.kz, "setCaseKey", session.__setitem__)
    monkeypatch.setattr(owb.kz, "storeCaseKey", session.__setitem__)

    ws = openpyxl.load_workbook(owb.saveContributions())[name]
    header = [c.value for c in ws[1]]
    saved = ws[row + 2]  # header row, then 1-based rows
    assert saved[header.index("Roth conv fixed")].value is True
    assert saved[header.index("Roth conv")].value == 12_345.0
