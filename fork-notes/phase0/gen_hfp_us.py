"""Generate an empty HFP workbook for the Phase 0 baseline case.

Run from the repository root:

    uv run python fork-notes/phase0/gen_hfp_us.py                          # otherFiles/HFP_us.xlsx
    uv run python fork-notes/phase0/gen_hfp_us.py otherFiles/HFP_us_2027.xlsx

It refuses to overwrite an existing workbook, so a filled-in one is never lost.

Columns per person sheet (exact headers, order may vary; only 'year' is required):
  year, anticipated wages, other inc, net inv,
  taxable ctrb, 401k ctrb, IRA ctrb, Roth 401k ctrb, Roth IRA ctrb, HSA ctrb,
  Roth conv, Roth conv fixed, QCD, big-ticket items

'anticipated wages' must be net of all contribution columns.
'big-ticket items' is the only signed column (negative = outflow).
"""

import sys
from pathlib import Path

import openpyxl

NAMES = ("SpouseA", "SpouseB")  # must match basic_info.names in the case file
YEARS = list(range(2026, 2057))  # through the year both reach 92 (born 1964)
COLUMNS = [
    "year", "anticipated wages", "other inc", "net inv",
    "taxable ctrb", "401k ctrb", "IRA ctrb", "Roth 401k ctrb", "Roth IRA ctrb", "HSA ctrb",
    "Roth conv", "Roth conv fixed", "QCD", "big-ticket items",
]

out = Path(sys.argv[1] if len(sys.argv) > 1 else "otherFiles/HFP_us.xlsx")
if out.exists():
    sys.exit(f"{out} exists; not overwriting it. Pass another path or delete it first.")
out.parent.mkdir(parents=True, exist_ok=True)

wb = openpyxl.Workbook()
wb.remove(wb.active)
for name in NAMES:
    ws = wb.create_sheet(name)
    ws.append(COLUMNS)
    for year in YEARS:
        ws.append([year] + [0] * (len(COLUMNS) - 1))
wb.save(out)

print(f"Wrote {out}: sheets {wb.sheetnames}, years {YEARS[0]}-{YEARS[-1]} ({len(YEARS)} rows)")
print("Fill in: anticipated wages (net of contributions) for the working years;")
print("         big-ticket items, negative, for housing and long-term-care outflows.")
